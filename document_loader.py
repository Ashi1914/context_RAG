"""
DOCX Document Loader
Loads and extracts text from DOCX files for RAG applications.
"""

from pathlib import Path
from typing import List, Optional
from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter


def load_docx(file_path: str) -> List[Document]:
    """
    Load a DOCX file and return a list of Document objects.
    
    Args:
        file_path: Path to the DOCX file
        
    Returns:
        List of langchain Document objects with extracted text
    """
    try:
        from docx import Document as DocxDocument
    except ImportError:
        raise ImportError(
            "python-docx is required to load DOCX files. "
            "Install it with: pip install python-docx"
        )
    
    file_path = Path(file_path)
    
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")
    
    if file_path.suffix.lower() != ".docx":
        raise ValueError(f"File must be a DOCX file, got: {file_path.suffix}")
    
    try:
        doc = DocxDocument(str(file_path))
    except Exception as e:
        raise ValueError(f"Failed to read DOCX file: {e}")
    
    documents = []
    current_heading = "General"
    
    # Extract text from paragraphs
    for para in doc.paragraphs:
        # Detect headings (standard styles)
        if para.style.name.startswith('Heading'):
            current_heading = para.text.strip()
            
        if para.text.strip():
            doc_obj = Document(
                page_content=para.text,
                metadata={
                    "source": str(file_path),
                    "file_name": file_path.name,
                    "file_type": "docx",
                    "section": current_heading
                }
            )
            documents.append(doc_obj)
    
    # Extract text from tables
    for table_idx, table in enumerate(doc.tables):
        for row_idx, row in enumerate(table.rows):
            row_text = " | ".join(
                cell.text.strip() for cell in row.cells
            )
            if row_text.strip():
                doc_obj = Document(
                    page_content=row_text,
                    metadata={
                        "source": str(file_path),
                        "file_name": file_path.name,
                        "file_type": "docx",
                        "table": table_idx,
                        "row": row_idx
                    }
                )
                documents.append(doc_obj)
    
    return documents


def load_docx_directory(directory_path: str, pattern: str = "*.docx") -> List[Document]:
    """
    Load all DOCX files from a directory.
    
    Args:
        directory_path: Path to the directory containing DOCX files
        pattern: File pattern to match (default: "*.docx")
        
    Returns:
        List of langchain Document objects from all matching files
    """
    directory = Path(directory_path)
    
    if not directory.is_dir():
        raise ValueError(f"Directory not found: {directory}")
    
    documents = []
    docx_files = list(directory.glob(pattern))
    
    if not docx_files:
        print(f"No DOCX files found in {directory}")
        return documents
    
    for file_path in docx_files:
        try:
            docs = load_docx(str(file_path))
            documents.extend(docs)
            print(f"[OK] Loaded {len(docs)} chunks from {file_path.name}")
        except Exception as e:
            print(f"[ERROR] Error loading {file_path.name}: {e}")
    
    return documents


def load_and_split(file_path: str, chunk_size: int = 2000, chunk_overlap: int = 200) -> List[Document]:
    """
    Load a PDF or DOCX file and split it into chunks for RAG applications.
    
    Args:
        file_path: Path to the file
        chunk_size: Size of each chunk
        chunk_overlap: Overlap between chunks
        
    Returns:
        List of langchain Document objects split into chunks
    """
    file_path_obj = Path(file_path)
    extension = file_path_obj.suffix.lower()
    
    if extension == ".pdf":
        loader = PyPDFLoader(file_path)
        docs = loader.load()
    elif extension == ".docx":
        docs = load_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {extension}")

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap
    )

    return splitter.split_documents(docs)


if __name__ == "__main__":
    # Example usage
    import sys
    
    if len(sys.argv) > 1:
        file_path = sys.argv[1]
        docs = load_docx(file_path)
        print(f"Loaded {len(docs)} documents from {file_path}")
        for i, doc in enumerate(docs[:3]):
            print(f"\nDocument {i+1}:")
            print(doc.page_content[:200])
    else:
        print("Usage: python docx_loader.py <path_to_docx_file>")
